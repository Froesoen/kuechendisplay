"""Liest Excel/CSV in eine Liste von Roh-Zeilen (dict mit logischen Feldnamen).

Zwei Layouts werden erkannt:
  * Mit Kopfzeile: Spaltennamen werden ueber ALIASES zugeordnet.
  * Ohne Kopfzeile: festes Layout der Adressliste (POSITIONAL).
Komplett leere Zeilen bleiben erhalten (Leerstring-Dict), weil sie Haushaltsbloecke trennen.
Nicht zugeordnete Spalten der Kopfzeile stehen in rows[0]["_unmapped"] (der Mapper meldet sie als Hinweis).
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

FIELDS = [
    "last_name", "first_name", "birthday", "street", "postal_code", "city",
    "phone_home", "email", "phone_mobile", "flag", "note", "category",
]
# Layout der vorhandenen Adressliste (Adressen_*.xlsx), ohne Kopfzeile.
# Spalte 12 (Kategorie) ist optional und fehlt in aelteren Tabellen.
POSITIONAL = FIELDS

ALIASES: dict[str, list[str]] = {
    "last_name": ["nachname", "name", "familienname", "last_name", "lastname"],
    "first_name": ["vorname", "vornamen", "first_name", "firstname"],
    "birthday": ["geburtstag", "geburtsdatum", "geburtstage", "birthday", "bday"],
    "street": ["strasse", "straße", "strasse und hausnummer", "street", "adresse"],
    "postal_code": ["plz", "postleitzahl", "postal_code", "zip"],
    "city": ["ort", "stadt", "wohnort", "city"],
    "phone_home": ["telefon", "festnetz", "festnetznummer", "telefonnummer", "tel", "phone", "phone_home"],
    "email": ["email", "e-mail", "mail", "emailadresse", "e-mailadresse", "mailadresse"],
    "phone_mobile": ["mobil", "mobilnummer", "handy", "handynummer", "mobiltelefon", "mobile", "phone_mobile"],
    "flag": ["flag", "kennzeichen"],
    "note": ["notiz", "notizen", "bemerkung", "bemerkungen", "note", "notes"],
    "category": ["kategorie", "kategorien", "category", "tag", "tags", "gruppe"],
}


def _clean(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_xlsx(path: Path, sheet: str | None) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _read_csv(path: Path) -> list[list[Any]]:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[_clean(c) for c in row] for row in csv.reader(text.splitlines(), dialect)]


def _norm_header(s: str) -> str:
    return str(s).strip().lower().replace(" ", "").replace("_", "").replace("-", "")


def _header_map(first_row: list[Any]) -> dict[int, str] | None:
    """Gibt {spaltenindex: feld} zurueck, wenn die Zeile wie eine Kopfzeile aussieht."""
    lookup = {_norm_header(a): f for f, al in ALIASES.items() for a in al}
    mapping = {i: lookup[_norm_header(c)] for i, c in enumerate(first_row)
               if isinstance(c, str) and _norm_header(c) in lookup}
    names = set(mapping.values())
    return mapping if {"first_name", "last_name"} & names and len(mapping) >= 2 else None


def sheet_names(path: str | Path) -> list[str]:
    """Tabellenblaetter einer xlsx-Datei (leer bei CSV)."""
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return []
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    return list(wb.sheetnames)


def read_table(path: str | Path, sheet: str | None = None) -> list[dict[str, Any]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        rows = _read_xlsx(path, sheet)
    elif suffix in (".csv", ".txt"):
        rows = _read_csv(path)
    else:
        raise ValueError(f"Nicht unterstuetztes Dateiformat: {suffix}")
    if not rows:
        return []

    header = _header_map(rows[0])
    unmapped = ""
    if header:
        body, mapping = rows[1:], header
        first_line = 2
        unmapped = ", ".join(str(c) for i, c in enumerate(rows[0]) if str(c).strip() and i not in header)
    else:
        body, mapping = rows, dict(enumerate(POSITIONAL))
        first_line = 1

    out: list[dict[str, Any]] = []
    for n, row in enumerate(body, start=first_line):
        rec: dict[str, Any] = {f: "" for f in FIELDS}
        for idx, field in mapping.items():
            if idx < len(row):
                rec[field] = row[idx]
        rec["_line"] = n
        out.append(rec)
    if out and unmapped:
        out[0]["_unmapped"] = unmapped
    return out
